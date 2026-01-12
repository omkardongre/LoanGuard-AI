"""
PowerPoint Report Generation for LoanGuard AI Platform.

Uses python-pptx to generate professional Risk Committee presentations.
V10.1 Architecture: P1 Feature - PowerPoint Export for Risk Committee

Production-level implementation with real data integration.
"""

import io
import logging
from datetime import datetime, date
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData

logger = logging.getLogger(__name__)


# Brand colors matching pdf_report_generator.py
COLORS = {
    "primary": RGBColor(30, 64, 175),     # #1e40af - Blue
    "success": RGBColor(21, 128, 61),     # #15803d - Green
    "warning": RGBColor(202, 138, 4),     # #ca8a04 - Amber
    "danger": RGBColor(220, 38, 38),      # #dc2626 - Red
    "muted": RGBColor(107, 114, 128),     # #6b7280 - Gray
    "white": RGBColor(255, 255, 255),
    "black": RGBColor(0, 0, 0),
    "background": RGBColor(248, 250, 252),  # #f8fafc
}


class RiskCommitteePresentation:
    """
    Generate Risk Committee PowerPoint Presentations.
    
    Production-level implementation for LMA Edge Hackathon.
    Generates professional board-ready presentations from real loan data.
    """
    
    # Slide dimensions (standard 16:9)
    SLIDE_WIDTH = Inches(13.333)
    SLIDE_HEIGHT = Inches(7.5)
    
    def __init__(self):
        """Initialize presentation generator."""
        self.prs = None
    
    def _create_presentation(self) -> Presentation:
        """Create new presentation with default settings."""
        prs = Presentation()
        prs.slide_width = self.SLIDE_WIDTH
        prs.slide_height = self.SLIDE_HEIGHT
        return prs
    
    def _get_status_color(self, status: str) -> RGBColor:
        """Get color for compliance status."""
        if not status:
            return COLORS["muted"]
        
        status_upper = status.upper()
        if status_upper in ("GREEN", "COMPLIANT", "ON_TRACK", "LOW", "VERIFIED", "PASS"):
            return COLORS["success"]
        elif status_upper in ("AMBER", "WARNING", "MEDIUM", "QUESTIONABLE"):
            return COLORS["warning"]
        elif status_upper in ("RED", "BREACH", "HIGH", "CRITICAL", "FAIL"):
            return COLORS["danger"]
        return COLORS["muted"]
    
    def _add_title_slide(
        self,
        prs: Presentation,
        title: str,
        subtitle: str,
    ) -> None:
        """Add title slide with LoanGuard AI branding."""
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)
        
        # Background color bar at top
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            self.SLIDE_WIDTH, Inches(1.5)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["primary"]
        shape.line.fill.background()
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.5),
            Inches(12), Inches(1.5)
        )
        title_frame = title_box.text_frame
        title_para = title_frame.paragraphs[0]
        title_para.text = title
        title_para.font.size = Pt(44)
        title_para.font.bold = True
        title_para.font.color.rgb = COLORS["primary"]
        title_para.alignment = PP_ALIGN.CENTER
        
        # Subtitle
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(4),
            Inches(12), Inches(0.8)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.text = subtitle
        subtitle_para.font.size = Pt(24)
        subtitle_para.font.color.rgb = COLORS["muted"]
        subtitle_para.alignment = PP_ALIGN.CENTER
        
        # Footer with branding
        self._add_footer(slide, "LoanGuard AI | Risk Committee Presentation")
    
    def _add_footer(self, slide, text: str) -> None:
        """Add footer to slide."""
        footer_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(6.8),
            Inches(12), Inches(0.4)
        )
        footer_frame = footer_box.text_frame
        footer_para = footer_frame.paragraphs[0]
        footer_para.text = f"{text} | Generated: {datetime.now().strftime('%Y-%m-%d')}"
        footer_para.font.size = Pt(10)
        footer_para.font.color.rgb = COLORS["muted"]
        footer_para.alignment = PP_ALIGN.CENTER
    
    def _add_section_header(self, slide, title: str) -> None:
        """Add section header to slide."""
        # Header bar
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            self.SLIDE_WIDTH, Inches(0.8)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["primary"]
        shape.line.fill.background()
        
        # Title in header
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.2),
            Inches(12), Inches(0.5)
        )
        title_frame = title_box.text_frame
        title_para = title_frame.paragraphs[0]
        title_para.text = title
        title_para.font.size = Pt(24)
        title_para.font.bold = True
        title_para.font.color.rgb = COLORS["white"]
    
    def _add_executive_summary_slide(
        self,
        prs: Presentation,
        summary: Dict[str, Any],
    ) -> None:
        """Add executive summary slide with key metrics."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "Executive Summary")
        
        # Key metrics grid
        metrics = [
            ("Total Loans", str(summary.get("total_loans", 0))),
            ("Total Exposure", f"${summary.get('total_exposure', 0):,.0f}"),
            ("Compliant", str(summary.get("loans_compliant", 0))),
            ("Warning", str(summary.get("loans_warning", 0))),
            ("Breach", str(summary.get("loans_breach", 0))),
            ("ESG Score", f"{summary.get('esg_average_score', 0):.1f}"),
        ]
        
        # Create metric boxes
        start_x = Inches(0.5)
        start_y = Inches(1.2)
        box_width = Inches(4)
        box_height = Inches(1.2)
        
        for i, (label, value) in enumerate(metrics):
            col = i % 3
            row = i // 3
            x = start_x + col * (box_width + Inches(0.3))
            y = start_y + row * (box_height + Inches(0.2))
            
            # Box shape
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                x, y, box_width, box_height
            )
            box.fill.solid()
            box.fill.fore_color.rgb = COLORS["background"]
            box.line.color.rgb = COLORS["muted"]
            
            # Value text
            value_box = slide.shapes.add_textbox(
                x + Inches(0.2), y + Inches(0.2),
                box_width - Inches(0.4), Inches(0.6)
            )
            value_frame = value_box.text_frame
            value_para = value_frame.paragraphs[0]
            value_para.text = value
            value_para.font.size = Pt(32)
            value_para.font.bold = True
            value_para.font.color.rgb = COLORS["primary"]
            value_para.alignment = PP_ALIGN.CENTER
            
            # Label text
            label_box = slide.shapes.add_textbox(
                x + Inches(0.2), y + Inches(0.7),
                box_width - Inches(0.4), Inches(0.4)
            )
            label_frame = label_box.text_frame
            label_para = label_frame.paragraphs[0]
            label_para.text = label
            label_para.font.size = Pt(14)
            label_para.font.color.rgb = COLORS["muted"]
            label_para.alignment = PP_ALIGN.CENTER
        
        # Overall status indicator
        overall_status = summary.get("overall_status", "GREEN")
        status_color = self._get_status_color(overall_status)
        
        status_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(4.5), Inches(4),
            Inches(4), Inches(1.5)
        )
        status_box.fill.solid()
        status_box.fill.fore_color.rgb = status_color
        status_box.line.fill.background()
        
        status_text = slide.shapes.add_textbox(
            Inches(4.5), Inches(4.3),
            Inches(4), Inches(1)
        )
        status_frame = status_text.text_frame
        status_para = status_frame.paragraphs[0]
        status_para.text = f"PORTFOLIO STATUS: {overall_status}"
        status_para.font.size = Pt(20)
        status_para.font.bold = True
        status_para.font.color.rgb = COLORS["white"]
        status_para.alignment = PP_ALIGN.CENTER
        
        self._add_footer(slide, "LoanGuard AI | Executive Summary")
    
    def _add_covenant_slide(
        self,
        prs: Presentation,
        covenants: List[Dict[str, Any]],
    ) -> None:
        """Add covenant compliance slide with table."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "Covenant Compliance Overview")
        
        if not covenants:
            # No data message
            no_data = slide.shapes.add_textbox(
                Inches(3), Inches(3),
                Inches(7), Inches(1)
            )
            no_data_frame = no_data.text_frame
            no_data_para = no_data_frame.paragraphs[0]
            no_data_para.text = "No covenant data available"
            no_data_para.font.size = Pt(24)
            no_data_para.font.color.rgb = COLORS["muted"]
            no_data_para.alignment = PP_ALIGN.CENTER
            self._add_footer(slide, "LoanGuard AI | Covenant Compliance")
            return
        
        # Summary counts
        compliant = sum(1 for c in covenants if c.get("status", "").upper() in ("GREEN", "COMPLIANT"))
        warning = sum(1 for c in covenants if c.get("status", "").upper() in ("AMBER", "WARNING"))
        breach = sum(1 for c in covenants if c.get("status", "").upper() in ("RED", "BREACH"))
        
        # Summary text
        summary_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.0),
            Inches(12), Inches(0.5)
        )
        summary_frame = summary_box.text_frame
        summary_para = summary_frame.paragraphs[0]
        summary_para.text = f"Total: {len(covenants)} | Compliant: {compliant} | Warning: {warning} | Breach: {breach}"
        summary_para.font.size = Pt(14)
        summary_para.font.color.rgb = COLORS["muted"]
        
        # Create table
        rows = min(len(covenants) + 1, 8)  # Header + up to 7 data rows
        cols = 5
        
        table_shape = slide.shapes.add_table(
            rows, cols,
            Inches(0.5), Inches(1.6),
            Inches(12), Inches(rows * 0.6)
        )
        table = table_shape.table
        
        # Set column widths
        table.columns[0].width = Inches(3)     # Covenant name
        table.columns[1].width = Inches(2)     # Threshold
        table.columns[2].width = Inches(2)     # Actual
        table.columns[3].width = Inches(2.5)   # Buffer
        table.columns[4].width = Inches(2.5)   # Status
        
        # Header row
        headers = ["Covenant", "Threshold", "Actual", "Buffer %", "Status"]
        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["primary"]
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(12)
            para.font.bold = True
            para.font.color.rgb = COLORS["white"]
            para.alignment = PP_ALIGN.CENTER
        
        # Data rows
        for row_idx, cov in enumerate(covenants[:7], 1):  # Limit to 7 rows
            status = cov.get("status", "N/A")
            status_color = self._get_status_color(status)
            
            row_data = [
                cov.get("name", cov.get("covenant_type", "N/A")),
                str(cov.get("threshold", "N/A")),
                str(cov.get("actual", cov.get("actual_value", "N/A"))),
                f"{cov.get('buffer_pct', 0):.1f}%",
                status,
            ]
            
            for col_idx, value in enumerate(row_data):
                cell = table.cell(row_idx, col_idx)
                cell.text = str(value)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(11)
                para.alignment = PP_ALIGN.CENTER
                
                # Color the status cell
                if col_idx == 4:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = status_color
                    para.font.color.rgb = COLORS["white"]
                    para.font.bold = True
        
        self._add_footer(slide, "LoanGuard AI | Covenant Compliance")
    
    def _add_ml_predictions_slide(
        self,
        prs: Presentation,
        predictions: Dict[str, Any],
    ) -> None:
        """Add ML predictions slide with risk factors."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "ML Breach Predictions")
        
        if not predictions:
            no_data = slide.shapes.add_textbox(
                Inches(3), Inches(3),
                Inches(7), Inches(1)
            )
            no_data_frame = no_data.text_frame
            no_data_para = no_data_frame.paragraphs[0]
            no_data_para.text = "No prediction data available"
            no_data_para.font.size = Pt(24)
            no_data_para.font.color.rgb = COLORS["muted"]
            no_data_para.alignment = PP_ALIGN.CENTER
            self._add_footer(slide, "LoanGuard AI | ML Predictions")
            return
        
        # Main probability
        prob = predictions.get("breach_probability", predictions.get("90_day_probability", 0))
        prob_pct = prob * 100 if prob <= 1 else prob
        
        risk_level = "CRITICAL" if prob_pct > 75 else "HIGH" if prob_pct > 50 else "MEDIUM" if prob_pct > 25 else "LOW"
        risk_color = self._get_status_color(risk_level)
        
        # Probability indicator
        prob_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5), Inches(1.2),
            Inches(5), Inches(2)
        )
        prob_box.fill.solid()
        prob_box.fill.fore_color.rgb = risk_color
        prob_box.line.fill.background()
        
        prob_text = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.5),
            Inches(5), Inches(1)
        )
        prob_frame = prob_text.text_frame
        prob_para = prob_frame.paragraphs[0]
        prob_para.text = f"{prob_pct:.1f}%"
        prob_para.font.size = Pt(48)
        prob_para.font.bold = True
        prob_para.font.color.rgb = COLORS["white"]
        prob_para.alignment = PP_ALIGN.CENTER
        
        prob_label = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.4),
            Inches(5), Inches(0.5)
        )
        prob_label_frame = prob_label.text_frame
        prob_label_para = prob_label_frame.paragraphs[0]
        prob_label_para.text = "90-Day Breach Probability"
        prob_label_para.font.size = Pt(14)
        prob_label_para.font.color.rgb = COLORS["white"]
        prob_label_para.alignment = PP_ALIGN.CENTER
        
        # Risk factors
        factors = predictions.get("top_risk_factors", predictions.get("risk_factors", []))
        
        if factors:
            factors_title = slide.shapes.add_textbox(
                Inches(6), Inches(1.2),
                Inches(6), Inches(0.5)
            )
            factors_title_frame = factors_title.text_frame
            factors_title_para = factors_title_frame.paragraphs[0]
            factors_title_para.text = "Top Risk Factors (SHAP)"
            factors_title_para.font.size = Pt(18)
            factors_title_para.font.bold = True
            factors_title_para.font.color.rgb = COLORS["primary"]
            
            # Factor list
            factors_box = slide.shapes.add_textbox(
                Inches(6), Inches(1.8),
                Inches(6.5), Inches(3)
            )
            factors_frame = factors_box.text_frame
            factors_frame.word_wrap = True
            
            for i, factor in enumerate(factors[:5]):
                if i > 0:
                    para = factors_frame.add_paragraph()
                else:
                    para = factors_frame.paragraphs[0]
                
                if isinstance(factor, dict):
                    factor_name = factor.get("feature", factor.get("name", "Unknown"))
                    factor_value = factor.get("importance", factor.get("value", 0))
                    para.text = f"• {factor_name}: {factor_value:.3f}"
                else:
                    para.text = f"• {factor}"
                
                para.font.size = Pt(14)
                para.font.color.rgb = COLORS["black"]
                para.space_after = Pt(8)
        
        # Model info - only show if data is available
        model_info = slide.shapes.add_textbox(
            Inches(0.5), Inches(3.5),
            Inches(5), Inches(1.5)
        )
        model_frame = model_info.text_frame
        
        model_name = predictions.get("model_type", predictions.get("model_name"))
        if model_name:
            model_para = model_frame.paragraphs[0]
            model_para.text = f"Model: {model_name}"
            model_para.font.size = Pt(12)
            model_para.font.color.rgb = COLORS["muted"]
        
        confidence = predictions.get("confidence")
        if confidence is not None:
            conf_para = model_frame.add_paragraph()
            conf_para.text = f"Confidence: {confidence * 100:.0f}%"
            conf_para.font.size = Pt(12)
            conf_para.font.color.rgb = COLORS["muted"]
        
        auc = predictions.get("auc", predictions.get("model_auc"))
        if auc is not None:
            auc_para = model_frame.add_paragraph()
            auc_para.text = f"Model AUC: {auc * 100:.2f}%"
            auc_para.font.size = Pt(12)
            auc_para.font.color.rgb = COLORS["muted"]
        
        self._add_footer(slide, "LoanGuard AI | ML Predictions (SHAP Explained)")
    
    def _add_esg_slide(
        self,
        prs: Presentation,
        esg_data: Dict[str, Any],
    ) -> None:
        """Add ESG/TLP compliance slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "ESG & TLP Compliance Status")
        
        if not esg_data:
            no_data = slide.shapes.add_textbox(
                Inches(3), Inches(3),
                Inches(7), Inches(1)
            )
            no_data_frame = no_data.text_frame
            no_data_para = no_data_frame.paragraphs[0]
            no_data_para.text = "No ESG data available"
            no_data_para.font.size = Pt(24)
            no_data_para.font.color.rgb = COLORS["muted"]
            no_data_para.alignment = PP_ALIGN.CENTER
            self._add_footer(slide, "LoanGuard AI | ESG Compliance")
            return
        
        # ESG Score
        esg_score = esg_data.get("overall_score", esg_data.get("esg_score", 0))
        esg_status = esg_data.get("overall_status", "N/A")
        status_color = self._get_status_color(esg_status)
        
        score_box = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(1), Inches(1.5),
            Inches(3), Inches(3)
        )
        score_box.fill.solid()
        score_box.fill.fore_color.rgb = status_color
        score_box.line.fill.background()
        
        score_text = slide.shapes.add_textbox(
            Inches(1), Inches(2.3),
            Inches(3), Inches(1)
        )
        score_frame = score_text.text_frame
        score_para = score_frame.paragraphs[0]
        score_para.text = f"{esg_score:.0f}"
        score_para.font.size = Pt(48)
        score_para.font.bold = True
        score_para.font.color.rgb = COLORS["white"]
        score_para.alignment = PP_ALIGN.CENTER
        
        score_label = slide.shapes.add_textbox(
            Inches(1), Inches(3.2),
            Inches(3), Inches(0.5)
        )
        score_label_frame = score_label.text_frame
        score_label_para = score_label_frame.paragraphs[0]
        score_label_para.text = "ESG Score"
        score_label_para.font.size = Pt(14)
        score_label_para.font.color.rgb = COLORS["white"]
        score_label_para.alignment = PP_ALIGN.CENTER
        
        # KPI summary
        kpis = esg_data.get("kpis", [])
        if kpis:
            kpi_title = slide.shapes.add_textbox(
                Inches(5), Inches(1.2),
                Inches(7), Inches(0.5)
            )
            kpi_title_frame = kpi_title.text_frame
            kpi_title_para = kpi_title_frame.paragraphs[0]
            kpi_title_para.text = "Sustainability KPIs"
            kpi_title_para.font.size = Pt(18)
            kpi_title_para.font.bold = True
            kpi_title_para.font.color.rgb = COLORS["primary"]
            
            # KPI table
            rows = min(len(kpis) + 1, 6)
            table_shape = slide.shapes.add_table(
                rows, 4,
                Inches(5), Inches(1.8),
                Inches(7.5), Inches(rows * 0.5)
            )
            table = table_shape.table
            
            table.columns[0].width = Inches(2.5)
            table.columns[1].width = Inches(1.5)
            table.columns[2].width = Inches(1.5)
            table.columns[3].width = Inches(2)
            
            headers = ["KPI", "Target", "Current", "Status"]
            for i, header in enumerate(headers):
                cell = table.cell(0, i)
                cell.text = header
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["primary"]
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(11)
                para.font.bold = True
                para.font.color.rgb = COLORS["white"]
            
            for row_idx, kpi in enumerate(kpis[:5], 1):
                row_data = [
                    kpi.get("name", "N/A"),
                    str(kpi.get("target", "N/A")),
                    str(kpi.get("current", "N/A")),
                    kpi.get("status", "N/A"),
                ]
                for col_idx, value in enumerate(row_data):
                    cell = table.cell(row_idx, col_idx)
                    cell.text = str(value)
                    para = cell.text_frame.paragraphs[0]
                    para.font.size = Pt(10)
        
        self._add_footer(slide, "LoanGuard AI | ESG Compliance")
    
    def _add_high_risk_loans_slide(
        self,
        prs: Presentation,
        loans: List[Dict[str, Any]],
    ) -> None:
        """Add high-risk loans slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "High-Risk Loans Requiring Attention")
        
        # Filter and sort loans by risk
        high_risk_loans = sorted(
            [l for l in loans if l.get("status", "").upper() in ("RED", "BREACH", "HIGH", "CRITICAL")],
            key=lambda x: x.get("breach_probability", x.get("risk_score", 0)),
            reverse=True
        )[:5]
        
        if not high_risk_loans:
            no_data = slide.shapes.add_textbox(
                Inches(3), Inches(3),
                Inches(7), Inches(1)
            )
            no_data_frame = no_data.text_frame
            no_data_para = no_data_frame.paragraphs[0]
            no_data_para.text = "No high-risk loans identified"
            no_data_para.font.size = Pt(24)
            no_data_para.font.color.rgb = COLORS["success"]
            no_data_para.alignment = PP_ALIGN.CENTER
            self._add_footer(slide, "LoanGuard AI | High-Risk Loans")
            return
        
        # Create table
        rows = len(high_risk_loans) + 1
        table_shape = slide.shapes.add_table(
            rows, 5,
            Inches(0.5), Inches(1.2),
            Inches(12), Inches(rows * 0.7)
        )
        table = table_shape.table
        
        table.columns[0].width = Inches(2)
        table.columns[1].width = Inches(3)
        table.columns[2].width = Inches(2.5)
        table.columns[3].width = Inches(2)
        table.columns[4].width = Inches(2.5)
        
        headers = ["Loan ID", "Borrower", "Exposure", "Breach Prob", "Status"]
        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["danger"]
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(12)
            para.font.bold = True
            para.font.color.rgb = COLORS["white"]
        
        for row_idx, loan in enumerate(high_risk_loans, 1):
            prob = loan.get("breach_probability", loan.get("risk_score", 0))
            row_data = [
                loan.get("loan_id", "N/A"),
                loan.get("borrower_name", "N/A")[:25],
                f"${loan.get('facility_amount', 0):,.0f}",
                f"{prob * 100:.1f}%" if prob <= 1 else f"{prob:.1f}%",
                loan.get("status", "N/A"),
            ]
            for col_idx, value in enumerate(row_data):
                cell = table.cell(row_idx, col_idx)
                cell.text = str(value)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(11)
        
        self._add_footer(slide, "LoanGuard AI | High-Risk Loans")
    
    def _add_recommendations_slide(
        self,
        prs: Presentation,
        recommendations: List[str],
    ) -> None:
        """Add recommendations slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "Recommended Actions")
        
        if not recommendations:
            # No recommendations available - show message
            no_data = slide.shapes.add_textbox(
                Inches(3), Inches(3),
                Inches(7), Inches(1)
            )
            no_data_frame = no_data.text_frame
            no_data_para = no_data_frame.paragraphs[0]
            no_data_para.text = "No specific recommendations generated"
            no_data_para.font.size = Pt(24)
            no_data_para.font.color.rgb = COLORS["muted"]
            no_data_para.alignment = PP_ALIGN.CENTER
            self._add_footer(slide, "LoanGuard AI | Recommendations")
            return
        
        rec_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.5),
            Inches(12), Inches(5)
        )
        rec_frame = rec_box.text_frame
        rec_frame.word_wrap = True
        
        for i, rec in enumerate(recommendations[:8]):
            if i > 0:
                para = rec_frame.add_paragraph()
            else:
                para = rec_frame.paragraphs[0]
            
            para.text = f"• {rec}"
            para.font.size = Pt(18)
            para.font.color.rgb = COLORS["black"]
            para.space_after = Pt(12)
        
        self._add_footer(slide, "LoanGuard AI | Recommendations")
    
    def _add_appendix_slide(
        self,
        prs: Presentation,
    ) -> None:
        """Add appendix/methodology slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "Appendix: Methodology")
        
        methodology = [
            ("ML Models", "LightGBM ensemble with XGBoost validation"),
            ("Explainability", "SHAP (SHapley Additive exPlanations) for EU AI Act compliance"),
            ("Stress Testing", "Basel III NGFS climate scenarios"),
            ("ESG Framework", "EBA Pillar 3 ESG disclosure requirements"),
            ("Risk Committee", "Multi-agent AI with specialized analysis agents"),
            ("Data Source", "BigQuery integration with FRED macroeconomic data"),
        ]
        
        start_y = Inches(1.2)
        for i, (title, desc) in enumerate(methodology):
            y = start_y + i * Inches(0.9)
            
            title_box = slide.shapes.add_textbox(
                Inches(0.5), y,
                Inches(3), Inches(0.4)
            )
            title_frame = title_box.text_frame
            title_para = title_frame.paragraphs[0]
            title_para.text = title
            title_para.font.size = Pt(14)
            title_para.font.bold = True
            title_para.font.color.rgb = COLORS["primary"]
            
            desc_box = slide.shapes.add_textbox(
                Inches(3.5), y,
                Inches(9), Inches(0.4)
            )
            desc_frame = desc_box.text_frame
            desc_para = desc_frame.paragraphs[0]
            desc_para.text = desc
            desc_para.font.size = Pt(12)
            desc_para.font.color.rgb = COLORS["black"]
        
        self._add_footer(slide, "LoanGuard AI | Confidential")
    
    def _add_risk_distribution_slide(
        self,
        prs: Presentation,
        summary: Dict[str, Any],
    ) -> None:
        """Add portfolio risk distribution pie chart slide."""
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        
        self._add_section_header(slide, "Portfolio Risk Distribution")
        
        # Get distribution data
        compliant = summary.get("loans_compliant", 0)
        warning = summary.get("loans_warning", 0)
        breach = summary.get("loans_breach", 0)
        total = compliant + warning + breach
        
        if total == 0:
            # No data - show message
            no_data = slide.shapes.add_textbox(
                Inches(3), Inches(3),
                Inches(7), Inches(1)
            )
            no_data_frame = no_data.text_frame
            no_data_para = no_data_frame.paragraphs[0]
            no_data_para.text = "No portfolio data available"
            no_data_para.font.size = Pt(24)
            no_data_para.font.color.rgb = COLORS["muted"]
            no_data_para.alignment = PP_ALIGN.CENTER
            self._add_footer(slide, "LoanGuard AI | Risk Distribution")
            return
        
        # Create pie chart
        chart_data = CategoryChartData()
        chart_data.categories = ["Compliant", "Warning", "Breach"]
        chart_data.add_series("Distribution", (compliant, warning, breach))
        
        x, y = Inches(0.8), Inches(1.5)
        cx, cy = Inches(5.5), Inches(4.5)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, x, y, cx, cy, chart_data
        ).chart
        
        # Format chart
        chart.has_legend = True
        chart.legend.include_in_layout = False
        
        # Add summary metrics on the right
        metrics_box = slide.shapes.add_textbox(
            Inches(7), Inches(1.8),
            Inches(5), Inches(4)
        )
        metrics_frame = metrics_box.text_frame
        metrics_frame.word_wrap = True
        
        # Title
        title_para = metrics_frame.paragraphs[0]
        title_para.text = "Portfolio Summary"
        title_para.font.size = Pt(18)
        title_para.font.bold = True
        title_para.font.color.rgb = COLORS["primary"]
        title_para.space_after = Pt(15)
        
        # Metrics
        metrics = [
            ("Total Loans", str(total), COLORS["black"]),
            ("Compliant (GREEN)", f"{compliant} ({compliant/total*100:.1f}%)", COLORS["success"]),
            ("Warning (AMBER)", f"{warning} ({warning/total*100:.1f}%)", COLORS["warning"]),
            ("Breach (RED)", f"{breach} ({breach/total*100:.1f}%)", COLORS["danger"]),
        ]
        
        for label, value, color in metrics:
            p = metrics_frame.add_paragraph()
            p.text = f"{label}: {value}"
            p.font.size = Pt(14)
            p.font.color.rgb = color
            p.space_after = Pt(8)
        
        # Health indicator
        health_p = metrics_frame.add_paragraph()
        health_p.space_before = Pt(15)
        if breach == 0 and warning == 0:
            health_p.text = "✓ Portfolio Health: EXCELLENT"
            health_p.font.color.rgb = COLORS["success"]
        elif breach == 0:
            health_p.text = "⚠ Portfolio Health: GOOD"
            health_p.font.color.rgb = COLORS["warning"]
        else:
            health_p.text = "✗ Portfolio Health: NEEDS ATTENTION"
            health_p.font.color.rgb = COLORS["danger"]
        health_p.font.size = Pt(16)
        health_p.font.bold = True
        
        self._add_footer(slide, "LoanGuard AI | Risk Distribution")
    
    def generate_loan_presentation(
        self,
        loan_data: Dict[str, Any],
        covenants: List[Dict[str, Any]],
        predictions: Optional[Dict[str, Any]] = None,
        esg_data: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        """
        Generate loan-level Risk Committee presentation.
        
        Args:
            loan_data: Loan details from database
            covenants: List of covenant measurements
            predictions: ML breach predictions
            esg_data: ESG compliance data
            
        Returns:
            PPTX file as bytes
        """
        prs = self._create_presentation()
        
        # Title slide
        loan_id = loan_data.get("loan_id", "Unknown")
        borrower = loan_data.get("borrower_name", "Unknown Borrower")
        self._add_title_slide(
            prs,
            "Risk Committee Report",
            f"Loan: {loan_id} | {borrower} | {date.today().strftime('%B %Y')}"
        )
        
        # Executive summary (convert loan to summary format)
        summary = {
            "total_loans": 1,
            "total_exposure": loan_data.get("facility_amount", 0),
            "loans_compliant": 1 if loan_data.get("overall_status", "").upper() == "GREEN" else 0,
            "loans_warning": 1 if loan_data.get("overall_status", "").upper() == "AMBER" else 0,
            "loans_breach": 1 if loan_data.get("overall_status", "").upper() == "RED" else 0,
            "esg_average_score": esg_data.get("overall_score", 0) if esg_data else 0,
            "overall_status": loan_data.get("overall_status", "GREEN"),
        }
        self._add_executive_summary_slide(prs, summary)
        
        # Covenant compliance
        self._add_covenant_slide(prs, covenants)
        
        # ML predictions
        self._add_ml_predictions_slide(prs, predictions or {})
        
        # ESG status
        self._add_esg_slide(prs, esg_data or {})
        
        # Recommendations
        recommendations = []
        status = loan_data.get("overall_status", "").upper()
        if status == "RED":
            recommendations.extend([
                "IMMEDIATE: Schedule borrower meeting within 48 hours",
                "Review covenant waiver or amendment options",
                "Prepare credit committee escalation briefing",
            ])
        elif status == "AMBER":
            recommendations.extend([
                "Schedule quarterly review within 30 days",
                "Increase monitoring frequency to bi-weekly",
                "Request updated financial projections from borrower",
            ])
        else:
            recommendations.append("Continue standard monitoring procedures")
        
        self._add_recommendations_slide(prs, recommendations)
        
        # Appendix
        self._add_appendix_slide(prs)
        
        # Save to bytes
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        
        logger.info(f"Generated loan presentation for {loan_id}")
        
        return buffer.getvalue()
    
    def generate_portfolio_presentation(
        self,
        summary: Dict[str, Any],
        loans: List[Dict[str, Any]],
        concentration: Optional[Dict[str, Any]] = None,
        predictions: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        """
        Generate portfolio-level Risk Committee presentation.
        
        Args:
            summary: Portfolio summary statistics
            loans: List of loans in portfolio
            concentration: Concentration analysis
            predictions: Aggregate ML predictions
            
        Returns:
            PPTX file as bytes
        """
        prs = self._create_presentation()
        
        # Title slide
        self._add_title_slide(
            prs,
            "Portfolio Risk Committee Report",
            f"Loan Portfolio Analysis | {date.today().strftime('%B %Y')}"
        )
        
        # Executive summary
        self._add_executive_summary_slide(prs, summary)
        
        # Top covenants (aggregate from all loans)
        all_covenants = []
        for loan in loans[:10]:
            covenants = loan.get("covenants", [])
            for cov in covenants:
                cov["loan_id"] = loan.get("loan_id")
                all_covenants.append(cov)
        
        self._add_covenant_slide(prs, all_covenants[:7])
        
        # ML predictions (portfolio level)
        if predictions:
            self._add_ml_predictions_slide(prs, predictions)
        
        # High-risk loans
        self._add_high_risk_loans_slide(prs, loans)
        
        # ESG summary (aggregate)
        esg_summary = {
            "overall_score": summary.get("esg_average_score", 0),
            "overall_status": "GREEN" if summary.get("esg_average_score", 0) >= 70 else "AMBER",
            "kpis": [],
        }
        self._add_esg_slide(prs, esg_summary)
        
        # Risk Distribution Chart (pie chart)
        self._add_risk_distribution_slide(prs, summary)
        
        # Recommendations - only include if real issues found
        recommendations = []
        breach_count = summary.get("loans_breach", 0)
        warning_count = summary.get("loans_warning", 0)
        
        if breach_count > 0:
            recommendations.append(f"PRIORITY: Address {breach_count} loans in breach status")
        if warning_count > 0:
            recommendations.append(f"Monitor {warning_count} loans in warning status closely")
        # Only add generic message if there are real issues to track
        
        self._add_recommendations_slide(prs, recommendations)
        
        # Appendix
        self._add_appendix_slide(prs)
        
        # Save to bytes
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        
        logger.info(f"Generated portfolio presentation with {len(loans)} loans")
        
        return buffer.getvalue()


# Singleton instance
_pptx_generator: Optional[RiskCommitteePresentation] = None


def get_pptx_generator() -> RiskCommitteePresentation:
    """Get or create PPTX generator singleton."""
    global _pptx_generator
    if _pptx_generator is None:
        _pptx_generator = RiskCommitteePresentation()
    return _pptx_generator


def generate_loan_pptx(
    loan_data: Dict[str, Any],
    covenants: List[Dict[str, Any]],
    predictions: Optional[Dict[str, Any]] = None,
    esg_data: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Convenience function to generate loan presentation."""
    generator = get_pptx_generator()
    return generator.generate_loan_presentation(loan_data, covenants, predictions, esg_data)


def generate_portfolio_pptx(
    summary: Dict[str, Any],
    loans: List[Dict[str, Any]],
    concentration: Optional[Dict[str, Any]] = None,
    predictions: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Convenience function to generate portfolio presentation."""
    generator = get_pptx_generator()
    return generator.generate_portfolio_presentation(summary, loans, concentration, predictions)
